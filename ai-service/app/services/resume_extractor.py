import os
import re
import io
import fitz  # PyMuPDF
import docx

class ExtractionError(Exception):
    """Custom exception raised when resume text extraction fails."""
    pass

class ResumeExtractor:
    """
    Reusable text extraction engine for parsing PDF and DOCX resume files.
    Extracts text, cleans formatting artifacts, and calculates statistics.
    """

    @staticmethod
    def clean_text(raw_text: str) -> str:
        """
        Cleans raw extracted text by normalizing newlines, stripping trailing spaces,
        and removing non-printable/control characters.
        """
        if not raw_text:
            return ""

        # Remove control characters except standard whitespace/newlines
        text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', raw_text)
        
        # Replace 3 or more consecutive newlines with a double newline
        text = re.sub(r'\n{3,}', '\n\n', text)
        
        # Replace multiple horizontal spaces/tabs with a single space
        lines = [re.sub(r'[ \t]+', ' ', line).strip() for line in text.splitlines()]
        
        return "\n".join(lines).strip()

    @classmethod
    def extract_from_pdf(cls, file_input: bytes | str) -> dict:
        """
        Extract text from PDF using PyMuPDF (fitz).
        Supports file path (str) or binary byte stream (bytes).
        """
        try:
            if isinstance(file_input, bytes):
                doc = fitz.open(stream=file_input, filetype="pdf")
            else:
                doc = fitz.open(file_input)

            page_count = len(doc)
            full_text_list = []

            for page_num in range(page_count):
                page = doc.load_page(page_num)
                page_text = page.get_text("text")
                if page_text:
                    full_text_list.append(page_text)

            doc.close()

            raw_text = "\n\n".join(full_text_list)
            cleaned_text = cls.clean_text(raw_text)

            return {
                "file_type": "pdf",
                "extracted_text": cleaned_text,
                "character_count": len(cleaned_text),
                "word_count": len(cleaned_text.split()),
                "page_count": page_count,
            }
        except Exception as e:
            raise ExtractionError(f"Failed to extract text from PDF document: {str(e)}")

    @classmethod
    def extract_from_docx(cls, file_input: bytes | str) -> dict:
        """
        Extract text from DOCX using python-docx.
        Supports file path (str) or binary byte stream (bytes).
        """
        try:
            if isinstance(file_input, bytes):
                doc = docx.Document(io.BytesIO(file_input))
            else:
                doc = docx.Document(file_input)

            full_text_list = []

            # 1. Extract paragraph text
            for p in doc.paragraphs:
                if p.text.strip():
                    full_text_list.append(p.text)

            # 2. Extract table cell text
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        full_text_list.append(" | ".join(row_text))

            raw_text = "\n".join(full_text_list)
            cleaned_text = cls.clean_text(raw_text)

            return {
                "file_type": "docx",
                "extracted_text": cleaned_text,
                "character_count": len(cleaned_text),
                "word_count": len(cleaned_text.split()),
                "page_count": 1,  # DOCX does not expose discrete page objects natively
            }
        except Exception as e:
            raise ExtractionError(f"Failed to extract text from DOCX document: {str(e)}")

    @classmethod
    def extract(cls, file_input: bytes | str, filename: str = "") -> dict:
        """
        Auto-detects document format (PDF vs DOCX) based on filename or byte header
        and executes appropriate extraction strategy.
        """
        lower_filename = filename.lower()

        if lower_filename.endswith(".pdf"):
            return cls.extract_from_pdf(file_input)
        elif lower_filename.endswith(".docx"):
            return cls.extract_from_docx(file_input)

        # Fallback to byte signature inspection if filename extension not provided
        if isinstance(file_input, bytes):
            if file_input.startswith(b"%PDF"):
                return cls.extract_from_pdf(file_input)
            elif file_input.startswith(b"PK\x03\x04"):
                return cls.extract_from_docx(file_input)

        raise ExtractionError("Unsupported file format. Only PDF (.pdf) and Word (.docx) documents are supported.")
